from openpyxl import load_workbook

from django.contrib import admin, messages
from django.core.exceptions import ValidationError
from django.db import transaction
from django.shortcuts import redirect, render
from django.urls import path

from .forms import ExcelImportForm
from .models import Book, BorrowRequest


@admin.register(Book)
class BookAdmin(admin.ModelAdmin):

    list_display = (
        "title",
        "author",
        "total_copies",
        "available_copies",
        "borrowed_copies",
        "is_active",
    )

    list_filter = (
        "is_active",
    )

    search_fields = (
        "title",
        "author",
    )

    change_list_template = (
        "admin/library/book/change_list.html"
    )

    def get_urls(self):

        urls = super().get_urls()

        custom_urls = [
            path(
                "import-excel/",
                self.admin_site.admin_view(
                    self.import_excel
                ),
                name="library_book_import_excel",
            ),
        ]

        return custom_urls + urls

    def import_excel(self, request):

        if request.method == "POST":

            form = ExcelImportForm(
                request.POST,
                request.FILES
            )

            if form.is_valid():

                excel_file = form.cleaned_data["excel_file"]

                try:

                    workbook = load_workbook(
                        excel_file,
                        read_only=True,
                        data_only=True
                    )

                    sheet = workbook.active

                    rows = sheet.iter_rows(
                        values_only=True
                    )

                    headers = next(rows, None)

                    if not headers:

                        messages.error(
                            request,
                            "فایل اکسل خالی است."
                        )

                        return redirect(
                            "admin:library_book_import_excel"
                        )

                    headers = [
                        str(header).strip()
                        if header is not None
                        else ""
                        for header in headers
                    ]

                    required_headers = [
                        "title",
                        "author",
                        "description",
                        "total_copies",
                        "available_copies",
                        "is_active",
                    ]

                    missing_headers = [
                        header
                        for header in required_headers
                        if header not in headers
                    ]

                    if missing_headers:

                        messages.error(
                            request,
                            "این ستون‌ها در فایل اکسل وجود ندارند: "
                            + ", ".join(missing_headers)
                        )

                        return redirect(
                            "admin:library_book_import_excel"
                        )

                    data = []
                    errors = []

                    existing_titles = set(
                        Book.objects.values_list(
                            "title",
                            flat=True
                        )
                    )

                    file_titles = set()

                    for row_number, row in enumerate(
                        rows,
                        start=2
                    ):

                        row_data = dict(
                            zip(headers, row)
                        )

                        title = str(
                            row_data.get("title") or ""
                        ).strip()

                        if not title:

                            errors.append(
                                f"ردیف {row_number}: "
                                "عنوان کتاب خالی است."
                            )

                            continue

                        if title in file_titles:

                            errors.append(
                                f"ردیف {row_number}: "
                                f"کتاب «{title}» در همین فایل تکراری است."
                            )

                            continue

                        if title in existing_titles:

                            errors.append(
                                f"ردیف {row_number}: "
                                f"کتاب «{title}» قبلاً در سایت وجود دارد."
                            )

                            continue

                        file_titles.add(title)

                        author = str(
                            row_data.get("author") or ""
                        ).strip()

                        description = str(
                            row_data.get("description") or ""
                        ).strip()

                        try:

                            total_copies = int(
                                row_data.get("total_copies") or 1
                            )

                            available_value = (
                                row_data.get("available_copies")
                            )

                            if available_value in (None, ""):

                                available_copies = total_copies

                            else:

                                available_copies = int(
                                    available_value
                                )

                        except (ValueError, TypeError):

                            errors.append(
                                f"ردیف {row_number}: "
                                "تعداد نسخه‌ها باید عدد صحیح باشد."
                            )

                            continue

                        if total_copies < 1:

                            errors.append(
                                f"ردیف {row_number}: "
                                "تعداد کل نسخه‌ها باید حداقل ۱ باشد."
                            )

                            continue

                        if (
                            available_copies < 0
                            or available_copies > total_copies
                        ):

                            errors.append(
                                f"ردیف {row_number}: "
                                "تعداد نسخه موجود نامعتبر است."
                            )

                            continue

                        is_active_value = str(
                            row_data.get("is_active")
                            if row_data.get("is_active") is not None
                            else "1"
                        ).strip().lower()

                        is_active = is_active_value not in [
                            "0",
                            "false",
                            "no",
                            "خیر",
                            "غیرفعال",
                        ]

                        data.append(
                            Book(
                                title=title,
                                author=author,
                                description=description,
                                total_copies=total_copies,
                                available_copies=available_copies,
                                is_active=is_active,
                            )
                        )

                    if errors:

                        for error in errors[:10]:

                            messages.error(
                                request,
                                error
                            )

                        if len(errors) > 10:

                            messages.error(
                                request,
                                f"{len(errors) - 10} خطای دیگر نیز وجود دارد."
                            )

                        return redirect(
                            "admin:library_book_import_excel"
                        )

                    with transaction.atomic():

                        Book.objects.bulk_create(
                            data
                        )

                    messages.success(
                        request,
                        f"{len(data)} کتاب با موفقیت وارد شد."
                    )

                    return redirect(
                        "admin:library_book_changelist"
                    )

                except Exception as error:

                    messages.error(
                        request,
                        f"خطا در خواندن فایل اکسل: {error}"
                    )

                    return redirect(
                        "admin:library_book_import_excel"
                    )

        else:

            form = ExcelImportForm()

        return render(
            request,
            "admin/library/book/import_excel.html",
            {
                "form": form,
                "title": "ورود کتاب‌ها از اکسل",
            }
        )


@admin.register(BorrowRequest)
class BorrowRequestAdmin(admin.ModelAdmin):

    list_display = (
        "requester_name",
        "book",
        "phone",
        "status",
        "created_at",
    )

    list_filter = (
        "status",
    )

    search_fields = (
        "requester_name",
        "phone",
        "book__title",
    )

    def save_model(
        self,
        request,
        obj,
        form,
        change
    ):

        with transaction.atomic():

            old_status = None
            old_book_id = None

            if change:

                old_request = (
                    BorrowRequest.objects
                    .select_for_update()
                    .get(pk=obj.pk)
                )

                old_status = old_request.status
                old_book_id = old_request.book_id

            old_is_approved = (
                old_status == "approved"
            )

            new_is_approved = (
                obj.status == "approved"
            )

            # --------------------------------
            # approved -> approved
            # ولی کتاب عوض شده
            # --------------------------------

            if (
                change
                and old_is_approved
                and new_is_approved
                and old_book_id != obj.book_id
            ):

                old_book = (
                    Book.objects
                    .select_for_update()
                    .get(pk=old_book_id)
                )

                new_book = (
                    Book.objects
                    .select_for_update()
                    .get(pk=obj.book_id)
                )

                old_book.available_copies += 1

                old_book.save(
                    update_fields=[
                        "available_copies"
                    ]
                )

                if new_book.available_copies <= 0:

                    raise ValidationError(
                        "کتاب جدید نسخه موجود ندارد."
                    )

                new_book.available_copies -= 1

                new_book.save(
                    update_fields=[
                        "available_copies"
                    ]

                )

            # --------------------------------
            # pending -> approved
            # --------------------------------

            elif (
                not old_is_approved
                and new_is_approved
            ):

                book = (
                    Book.objects
                    .select_for_update()
                    .get(pk=obj.book_id)
                )

                if book.available_copies <= 0:

                    raise ValidationError(
                        "این کتاب دیگر نسخه موجود ندارد."
                    )

                book.available_copies -= 1

                book.save(
                    update_fields=[
                        "available_copies"
                    ]
                )

                messages.success(
                    request,
                    f"کتاب «{book.title}» به عنوان امانت داده‌شده ثبت شد."
                )

            # --------------------------------
            # approved -> returned
            # --------------------------------

            elif (
                old_is_approved
                and obj.status == "returned"
            ):

                book = (
                    Book.objects
                    .select_for_update()
                    .get(pk=obj.book_id)
                )

                if book.available_copies < book.total_copies:

                    book.available_copies += 1

                    book.save(
                        update_fields=[
                            "available_copies"
                        ]
                    )

                    messages.success(
                        request,
                        f"کتاب «{book.title}» به کتابخانه بازگردانده شد."
                    )

            # --------------------------------
            # approved -> rejected
            # --------------------------------

            elif (
                old_is_approved
                and obj.status == "rejected"
            ):

                book = (
                    Book.objects
                    .select_for_update()
                    .get(pk=obj.book_id)
                )

                if book.available_copies < book.total_copies:

                    book.available_copies += 1

                    book.save(
                        update_fields=[
                            "available_copies"
                        ]
                    )

            super().save_model(
                request,
                obj,
                form,
                change
            )

    def delete_model(
        self,
        request,
        obj
    ):

        with transaction.atomic():

            if obj.status == "approved":

                book = (
                    Book.objects
                    .select_for_update()
                    .get(pk=obj.book_id)
                )

                if book.available_copies < book.total_copies:

                    book.available_copies += 1

                    book.save(
                        update_fields=[
                            "available_copies"
                        ]
                    )

            super().delete_model(
                request,
                obj
            )